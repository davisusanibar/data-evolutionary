"""Slides 35-42: divisor de anexos, tres covers HTML y cuatro capturas del owner.

El pptx presenta; el PDF documenta. Cada cover (layout lienzo) muestra la
primera pagina del anexo tal como la pagino Chrome (preview PNG generado por
anexar.py --paginas) y una caja lateral con la fuente real y la remision al
anexo completo del PDF fusionado. Falla cerrado si el preview no existe:
los covers no se inventan; se derivan.

Los anexos A4/A5 (slides 39-42, T016) insertan las capturas del owner
(deck/anexos/): evidencia visual aportada de la sesion e2e — Flink UI y la
TUI de engram. Las capturas se copian tal cual, jamas se recortan ni se
retocan, y estan pineadas por sha256 en feature.spec.yaml; si falta el
archivo, el slide falla cerrado.

Orden del pipeline: build_estructura.py -> anexar.py --paginas -> llenar.py
-> render.py -> anexar.py --fusionar.
"""
from pathlib import Path

from pptx.util import Pt

from kcd_estilo import *

BASE = Path(__file__).resolve().parent
ANEXOS_DIR = BASE / "build" / "anexos"
CAPTURAS_DIR = BASE / "anexos"

# (clave, titulo del cover, ruta fuente legible, resumen de contenido)
ANEXOS = [
    ("a1",
     "EDAIOS: el sistema operativo completo (vista day-zero)",
     "edaiosv/docs/demos/edaios-operating-system.html",
     ["Constitución, 15 ADR, gates y catálogo de Core,",
      "tal como los proyecta el baseline day-zero."]),
    ("a2",
     "EDAIOS aplicado: la demo Flink hereda el gobierno de Core",
     "data-kcd2026/docs/edaios-operating-system-flink.html",
     ["Los 7 artículos y los gates de Core, lado a lado",
      "con el Constitution Check del pipeline Flink."]),
    ("a3",
     "El caso Flink completo: contrato, canary, evidencia",
     "data-kcd2026/docs/edaios-operating-system-flink-use-case.html",
     ["La página del caso: FR/SC, el .avsc como vista,",
      "el canary en rojo y la evidencia e2e del compose."]),
]

NOTAS_COVER = {
    "a1": "Este anexo es la vista day-zero completa del sistema operativo: la "
          "constitucion, los quince ADR y los gates de Core. No la recorremos "
          "en vivo; viaja paginada en el PDF, con texto seleccionable. El cover "
          "solo muestra la portada real, renderizada por Chrome desde el HTML.",
    "a2": "Aqui esta la vista aplicada: que hereda la demo Flink del gobierno "
          "de Core, articulo por articulo y gate por gate. Si alguien pregunta "
          "como se conecta la charla con el baseline, la respuesta esta en este "
          "anexo del PDF, no en una slide comprimida.",
    "a3": "Y el caso completo: contrato Avro, canary en rojo, evidencia e2e. Es "
          "la misma pagina publicada del caso, paginada tal cual. El deck conto "
          "la historia; este anexo permite auditarla linea por linea.",
}


def _cover(s, numero, clave, titulo_texto, ruta, resumen):
    png = ANEXOS_DIR / f"preview-{clave}.png"
    if not png.is_file():
        raise FileNotFoundError(
            f"falta el preview {png}: los covers de anexos se derivan del render "
            "de Chrome, no se inventan. Corre primero: "
            "cd deck && python3 anexar.py --paginas (tras build_estructura.py)")

    # Titulo a 20pt: los titulos de anexo son largos y a 25pt envolverian.
    caja(s, MARGEN, 0.75, 8.9, 0.45, titulo_texto, tamano=20, negrita=True)
    caja(s, MARGEN, 1.22, 8.9, 0.3,
         f"Anexo {numero} del PDF entregable — vista completa, texto seleccionable",
         tamano=12, color=TINTA_SUAVE)

    # Preview de la pagina 1, enmarcado. x >= FARO_X: su base (4.9) cae por
    # debajo de FARO_Y y la zona del faro no admite elementos a su izquierda.
    pic = s.shapes.add_picture(str(png), pulg(FARO_X), pulg(1.7), height=pulg(3.2))
    marco = bloque(s, pic.left / 914400, pic.top / 914400,
                   pic.width / 914400, pic.height / 914400,
                   relleno=None, borde=GRIS_BORDE, radio=False)
    marco.line.width = Pt(1.0)

    # Caja lateral: fuente real, condicion de derivado y remision al PDF.
    lado_x, lado_w = COL_DER, 4.3
    bloque(s, lado_x, 1.7, lado_w, 3.0, relleno=BLANCO, borde=GRIS_BORDE)
    caja(s, lado_x + 0.2, 1.88, lado_w - 0.4, 0.28, "Fuente", tamano=12, negrita=True)
    caja(s, lado_x + 0.2, 2.18, lado_w - 0.4, 0.35, ruta, tamano=9, fuente=MONO,
         color=TINTA_SUAVE)
    caja(s, lado_x + 0.2, 2.62, lado_w - 0.4, 0.55,
         ["Vista regenerable — derivado determinista;", "no se edita a mano."],
         tamano=11)
    caja(s, lado_x + 0.2, 3.30, lado_w - 0.4, 0.55,
         [f"Versión completa y paginada: anexo {numero}", "del PDF (kcd2026-completo.pdf)."],
         tamano=11, negrita=True)
    caja(s, lado_x + 0.2, 3.98, lado_w - 0.4, 0.6, resumen, tamano=10,
         color=TINTA_SUAVE)
    notas(s, NOTAS_COVER[clave])


def _captura(s, nombre, x, y, w):
    """Inserta una captura del owner tal cual (sin recortar ni retocar).

    add_picture con solo width preserva la proporcion real del archivo.
    Marco sutil GRIS_BORDE, mismo trato que los previews de los covers.
    Falla cerrado si el archivo no esta: las capturas son evidencia visual
    aportada por el owner; no se generan ni se sustituyen.
    """
    ruta = CAPTURAS_DIR / nombre
    if not ruta.is_file():
        raise FileNotFoundError(
            f"falta la captura {ruta}: las capturas del owner son evidencia "
            "visual aportada — se copian TAL CUAL (sin recortar ni retocar) a "
            "deck/anexos/ y estan pineadas por sha256 en feature.spec.yaml. "
            "Copia el archivo original del owner a esa carpeta y reejecuta el "
            "pipeline en orden (build_estructura.py -> anexar.py --paginas -> "
            "llenar.py -> render.py -> anexar.py --fusionar).")
    pic = s.shapes.add_picture(str(ruta), pulg(x), pulg(y), width=pulg(w))
    marco = bloque(s, x, y, pic.width / 914400, pic.height / 914400,
                   relleno=None, borde=GRIS_BORDE, radio=False)
    marco.line.width = Pt(1.0)
    return pic


# Kickers compartidos de los anexos de capturas (A4 = Flink UI, A5 = engram).
KICKER_A4 = "Anexo A4 · captura del owner durante la sesión e2e"
KICKER_A5 = "Anexo A5 · captura del owner de la TUI de engram"


def _capturas(slides):
    """Slides 39-42: anexos A4/A5 — las cuatro capturas del owner."""
    # --------------------------------------- 39 A4: job de la demo en Flink UI
    s = slides[38]
    titulo(s, "Flink en vivo: el job de la demo", kicker=KICKER_A4)
    pic = _captura(s, "Flink-Job-ID.jpeg", 2.7, 1.62, 6.1)   # 2.20:1 -> 2.78" alto
    # Caption ancho desde FARO_X: dos lineas sin envolver, por encima de la
    # banda azul inferior (~y=5.0).
    cy = (pic.top + pic.height) / 914400 + 0.13
    caja(s, FARO_X, cy, 7.9, 0.42,
         [("Mismo job que la evidencia e2e: ",
           "608ee13c… = job_id de sc-002-004-e2e-compose.json (spec 001)."),
          ("Operadores nombrados por FR: ",
           "FR-001 lectura Avro / FR-004 descarte → FR-002 revenue / FR-005 escritura.")],
         tamano=10, tamano_prefijo=10, color=TINTA_SUAVE)
    notas(s, "Esta captura la tomo el owner durante la sesion e2e. El Job ID "
             "coincide con el que registra la evidencia sc-002-004: no es un job "
             "de utileria, es el mismo que produjo la ventana de 150.0. Y el "
             "grafo habla el idioma de la spec: cada operador lleva el nombre "
             "del FR que implementa.")

    # ----------------------------------------- 40 A4: overview del dashboard
    s = slides[39]
    titulo(s, "Flink: el overview honesto", kicker=KICKER_A4)
    caja(s, MARGEN, 1.62, 8.9, 0.42,
         [("“Failed: 1” no es ruido: ",
           "es el primer intento del job, muerto por la KryoException del slide 29."),
          ("El dashboard registró la historia que el deck cuenta: ",
           "1 job RUNNING, 3 slots, 1 TaskManager… y su cicatriz.")],
         tamano=10.5, tamano_prefijo=10.5, color=TINTA_SUAVE)
    _captura(s, "Flink-Job.jpeg", MARGEN, 2.25, 8.9)         # 4.71:1 -> 1.89" alto
    notas(s, "El overview no se maquillo: el contador Failed en 1 es el primer "
             "intento del job, el que murio por la KryoException que contamos en "
             "el slide 29. Dejarlo a la vista es parte de la tesis: la evidencia "
             "honesta incluye los intentos fallidos, no solo el RUNNING final.")

    # -------------------------------------- 41 A5: memoria de la sesion (TUI)
    s = slides[40]
    titulo(s, "engram: la memoria de la sesión", kicker=KICKER_A5)
    caja(s, MARGEN, 1.75, 2.95, 2.3,
         ["Ocho observaciones guardadas por la sesión de construcción: el "
          "gobierno demostrándose, el canary en rojo, los bugs y las decisiones.",
          "",
          "Memoria persistente: el contexto sobrevive a la sesión que lo "
          "escribió."],
         tamano=11, color=TINTA)
    _captura(s, "Engram-Logs.png", 3.8, 1.62, 5.0)           # 1.72:1 -> 2.91" alto
    notas(s, "engram guardo ocho observaciones durante la construccion de la "
             "demo: decisiones, bugs, el canary en rojo, el gobierno "
             "demostrandose. Es la memoria persistente de la sesion: si la "
             "charla o la demo se reconstruyen, este contexto no se pierde.")

    # --------------------------------------- 42 A5: la tesis como observacion
    s = slides[41]
    titulo(s, "engram: la tesis, guardada", kicker=KICKER_A5)
    pic = _captura(s, "Engram-Logs-ID.png", 2.85, 1.62, 5.95)  # 2.08:1 -> 2.86" alto
    cy = (pic.top + pic.height) / 914400 + 0.12
    caja(s, 2.85, cy, 5.95, 0.4,
         [("Observación #7: ",
           "“KCD 2026: tres bugs que los tests verdes no atraparon”."),
          ("La tesis de la charla, registrada como memoria ",
           "con su What / Why / Learned.")],
         tamano=10, tamano_prefijo=10, color=TINTA_SUAVE)
    notas(s, "Y la observacion 7 es la tesis de la charla guardada como memoria: "
             "tres bugs que los tests verdes no atraparon, con su What, Why y "
             "Learned. La charla termina donde empezo: el conocimiento primero, "
             "tambien como memoria de la sesion que la construyo.")


def construir(slides):
    # ------------------------------------------------------ 35 divisor: anexos
    s = slides[34]
    divisor(s, "05", ["Anexos"], "Las vistas completas viajan en el PDF")
    notas(s, "Lo que sigue no se presenta: se entrega. Tres vistas HTML "
             "completas, paginadas por Chrome con texto seleccionable, van "
             "fusionadas al final del PDF; las slides solo muestran su portada. "
             "Y cierran los anexos A4 y A5: cuatro capturas del owner de la "
             "sesion e2e — la Flink UI y la memoria de engram, tal cual.")

    # ------------------------------------------------- 36..38 covers de anexos
    for i, (clave, titulo_texto, ruta, resumen) in enumerate(ANEXOS):
        _cover(slides[35 + i], f"A{i + 1}", clave, titulo_texto, ruta, resumen)

    # -------------------------------- 39..42 anexos A4/A5: capturas del owner
    _capturas(slides)
