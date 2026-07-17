#!/usr/bin/env python3
"""Fase 3: llena los 43 slides. El fondo y el layout ya vienen del template.

Orden del pipeline completo (los covers de anexos leen los previews que
genera anexar.py --paginas, y build_estructura borra build/):

    python3 build_estructura.py      # 1. estructura de 43 slides
    python3 anexar.py --paginas      # 2. HTML -> aN.pdf + preview-aN.png
    python3 llenar.py                # 3. este script
    python3 render.py                # 4. pptx -> kcd2026.pdf + PNGs
    python3 anexar.py --fusionar     # 5. deck + anexos -> kcd2026-completo.pdf
"""
from pptx import Presentation
from pptx.util import Emu, Pt
import contenido_1, contenido_diagramas, contenido_2, contenido_anexos
from kcd_estilo import pulg

p = Presentation("build/kcd2026.pptx")
slides = list(p.slides)
assert len(slides) == 43, f"esperaba 43, hay {len(slides)}"

contenido_1.construir(slides)
contenido_diagramas.construir(slides)
contenido_2.construir(slides)
contenido_anexos.construir(slides)

# Pase de correccion: el cuadro de titulo del layout de bullets mide 3.3" y
# envuelve cualquier titulo de mas de dos palabras, chocando con la lista.
# Se ensancha al ancho util de la columna izquierda.
ajustados = 0
for s in slides:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        x, y, w = sh.left, sh.top, sh.width
        if abs(x - pulg(0.6)) < pulg(0.08) and abs(y - pulg(1.2)) < pulg(0.15) and w < pulg(4.0):
            sh.width = pulg(5.4)
            sh.text_frame.word_wrap = True
            ajustados += 1

p.save("build/kcd2026.pptx")
print(f"contenido aplicado a 43 slides; {ajustados} titulos ensanchados")
