"""Helpers de estilo para el deck KCD Lima 2026.

Paleta y tipografia tomadas del template oficial, no inventadas:
azul KCD sobre gris claro, Arial. El lienzo es de 10 x 5.62 pulgadas.

Sistema global: los contenido_*.py solo consumen estos helpers; aqui viven
la reticula, la paleta, la tarjeta unica, el chrome de log, el conector
unico y los patrones de titulo, divisor y statement.
"""
import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

AZUL = RGBColor(0x0F, 0x87, 0xFF)
AZUL_OSCURO = RGBColor(0x00, 0x4C, 0xA8)
GRIS_BG = RGBColor(0xEF, 0xEF, 0xEF)
TINTA = RGBColor(0x26, 0x26, 0x26)
TINTA_SUAVE = RGBColor(0x5A, 0x5A, 0x5A)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
ROJO = RGBColor(0xC0, 0x2F, 0x2F)
VERDE = RGBColor(0x1E, 0x7A, 0x3C)

# Conector unico: gris neutro 1.5pt para secuencias; AZUL/VERDE/AZUL_OSCURO
# a 1.75pt para flujo semantico. Borde neutro de tarjeta en GRIS_BORDE.
GRIS_FLECHA = RGBColor(0x8A, 0x8A, 0x8A)
GRIS_BORDE = RGBColor(0xBB, 0xBB, 0xBB)

# Panel de log oscuro. ROJO_PANEL y VERDE_PANEL son las variantes de
# contraste de ROJO/VERDE para fondo oscuro (fail/ok legibles en el panel).
PANEL_OSCURO = RGBColor(0x1B, 0x1B, 0x1F)
PANEL_CHROME = RGBColor(0x2A, 0x2A, 0x30)
TEXTO_PANEL = RGBColor(0xD8, 0xD8, 0xDC)
TEXTO_CHROME = RGBColor(0xA0, 0xA0, 0xA8)
ROJO_PANEL = RGBColor(0xFF, 0x6B, 0x6B)
VERDE_PANEL = RGBColor(0x7A, 0xD1, 0x7A)

FUENTE = "Arial"
MONO = "Courier New"

ANCHO = 10.0
ALTO = 5.625

# Reticula: margen exterior >= 0.5"; x=0.55 en lienzo claro; MARGEN_AZUL
# alineado al glifo del titulo del layout de bullets para toda caja/log de
# ancho completo en slides azules; columna derecha siempre en x=5.15.
MARGEN = 0.55
MARGEN_AZUL = 0.68
COL_DER = 5.15

# Zona de exclusion del faro en el layout 4 (lienzo): ningun elemento
# generado con x < FARO_X puede tener y + alto > FARO_Y.
FARO_X = 1.45
FARO_Y = 4.2


def pulg(v):
    return Emu(int(v * 914400))


def texto_de(shape):
    return shape.text_frame.text.strip() if shape.has_text_frame else ""


def buscar(slide, fragmento):
    """Encuentra un cuadro de texto por un fragmento de su contenido."""
    for sh in slide.shapes:
        if sh.has_text_frame and fragmento.lower() in sh.text_frame.text.lower():
            return sh
    return None


def quitar(slide, fragmento):
    """Elimina un placeholder que no aplica (p.ej. lista numerada del layout)."""
    sh = buscar(slide, fragmento)
    if sh is not None:
        sh._element.getparent().remove(sh._element)


def reemplazar(shape, lineas, tamano=None, negrita=None, color=None):
    """Reemplaza el texto conservando el formato del primer run.

    Asignar text_frame.text colapsa el parrafo a un run sin estilo. Aqui se
    conserva el run original y se CLONA su rPr en cada linea nueva: sin eso, los
    runs anadidos pierden tamano y peso y caen al default del layout.
    Al final se normaliza space_before de los parrafos clonados para que todas
    las lineas compartan un solo ritmo vertical.
    """
    if isinstance(lineas, str):
        lineas = [lineas]
    tf = shape.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()

    modelo_rpr = p0.runs[0]._r.find(qn("a:rPr"))
    modelo_ppr = p0._p.find(qn("a:pPr"))

    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)

    def aplicar(run, txt):
        run.text = txt
        f = run.font
        f.name = FUENTE
        if tamano:
            f.size = Pt(tamano)
        if negrita is not None:
            f.bold = negrita
        if color is not None:
            f.color.rgb = color

    aplicar(p0.runs[0], lineas[0])
    for linea in lineas[1:]:
        p = tf.add_paragraph()
        if modelo_ppr is not None:
            p._p.insert(0, copy.deepcopy(modelo_ppr))
        r = p.add_run()
        if modelo_rpr is not None:
            nuevo = copy.deepcopy(modelo_rpr)
            viejo = r._r.find(qn("a:rPr"))
            if viejo is not None:
                r._r.remove(viejo)
            r._r.insert(0, nuevo)
        aplicar(r, linea)

    # Un solo ritmo vertical: todos los parrafos con el space_before de p0.
    sb = p0.space_before
    for p in list(tf.paragraphs)[1:]:
        if sb is not None:
            p.space_before = sb
        else:
            ppr = p._p.find(qn("a:pPr"))
            if ppr is not None:
                spc = ppr.find(qn("a:spcBef"))
                if spc is not None:
                    ppr.remove(spc)
    return shape


def caja(slide, x, y, w, h, texto="", tamano=14, negrita=False, color=None,
         alineacion=PP_ALIGN.LEFT, fuente=FUENTE, interlineado=None,
         anclaje=MSO_ANCHOR.TOP, negrita_lineas=None, tamano_prefijo=None):
    """Cuadro de texto simple.

    negrita_lineas: indices de lineas que van en bold (encabezados intra-caja)
    sin multiplicar cajas. Una linea puede ser una tupla (prefijo, resto): el
    prefijo va en bold (a tamano_prefijo si se da) y el resto en regular.
    """
    tb = slide.shapes.add_textbox(pulg(x), pulg(y), pulg(w), pulg(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anclaje
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lineas = texto.split("\n") if isinstance(texto, str) else texto
    tinta = color if color is not None else TINTA
    for i, linea in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alineacion
        if interlineado:
            p.line_spacing = interlineado
        if isinstance(linea, tuple):
            pref, resto = linea
            partes = [(pref, True, tamano_prefijo or tamano), (resto, False, tamano)]
        else:
            en_bold = negrita or bool(negrita_lineas and i in negrita_lineas)
            partes = [(linea, en_bold, tamano)]
        for txt, es_bold, tam in partes:
            r = p.add_run()
            r.text = txt
            r.font.name = fuente
            r.font.size = Pt(tam)
            r.font.bold = es_bold
            r.font.color.rgb = tinta
    return tb


def bloque(slide, x, y, w, h, relleno=None, borde=None, radio=True, transparencia=None):
    """Tarjeta unica del deck: radio 0.08, sin sombra, borde 1.25pt.

    Convencion: blanco+borde AZUL = contenido; blanco+GRIS_BORDE = neutro;
    relleno AZUL/VERDE = protagonista/gate.
    """
    forma = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radio else MSO_SHAPE.RECTANGLE,
        pulg(x), pulg(y), pulg(w), pulg(h))
    if radio:
        forma.adjustments[0] = 0.08
    if relleno is None:
        forma.fill.background()
    else:
        forma.fill.solid()
        forma.fill.fore_color.rgb = relleno
        if transparencia is not None:
            forma.fill.transparency = transparencia
    if borde is None:
        forma.line.fill.background()
    else:
        forma.line.color.rgb = borde
        forma.line.width = Pt(1.25)
    forma.shadow.inherit = False
    if forma.has_text_frame:
        forma.text_frame.word_wrap = True
        forma.text_frame.margin_left = pulg(0.08)
        forma.text_frame.margin_right = pulg(0.08)
    return forma


def rotular(forma, lineas, tamano=12, negrita=True, color=None):
    tf = forma.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(lineas, str):
        lineas = [lineas]
    for i, linea in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = linea
        r.font.name = FUENTE
        # Piso tipografico absoluto: 9pt tambien en sub-rotulos.
        r.font.size = Pt(tamano if i == 0 else max(tamano - 2, 9))
        r.font.bold = negrita if i == 0 else False
        r.font.color.rgb = color if color is not None else TINTA
    return forma


def flecha(slide, x1, y1, x2, y2, color=None, grosor=1.75, punteado=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pulg(x1), pulg(y1), pulg(x2), pulg(y2))
    c.line.color.rgb = color if color is not None else AZUL
    c.line.width = Pt(grosor)
    linea = c.line._get_or_add_ln()
    if punteado:
        guiones = etree.SubElement(linea, qn("a:prstDash"))
        guiones.set("val", "dash")
    cabeza = etree.SubElement(linea, qn("a:tailEnd"))
    cabeza.set("type", "triangle")
    cabeza.set("w", "med")
    cabeza.set("len", "med")
    return c


def log(slide, x, y, w, lineas, tamano=9, resaltar=None, titulo=None, alto_min=None):
    """Bloque de log real, en monoespaciada sobre panel oscuro.

    El TEXTO de los logs no se toca jamas: son transcripciones. El alto se
    deriva del cuerpo (nada de banda muerta) y `titulo` dibuja una franja de
    chrome de 0.24" con el nombre del archivo o comando: chrome informativo
    dentro de la tarjeta, nunca vacio, no una franja decorativa.
    Piso de mono: 9pt (9.5pt en paneles anchos >= 8.8").
    alto_min: alinea el borde inferior con un panel hermano mas alto sin
    tocar el texto (el cuerpo queda anclado arriba).
    """
    chrome = 0.24 if titulo else 0.0
    alto = len(lineas) * (tamano * 1.35 / 72) + 0.26 + chrome
    if alto_min is not None:
        alto = max(alto, alto_min)
    panel = bloque(slide, x, y, w, alto, relleno=PANEL_OSCURO, radio=False)
    if titulo:
        bloque(slide, x, y, w, 0.24, relleno=PANEL_CHROME, radio=False)
        et = slide.shapes.add_textbox(pulg(x + 0.12), pulg(y + 0.025), pulg(w - 0.24), pulg(0.2))
        etf = et.text_frame
        etf.word_wrap = False
        etf.margin_left = etf.margin_right = etf.margin_top = etf.margin_bottom = 0
        rp = etf.paragraphs[0].add_run()
        rp.text = titulo
        rp.font.name = MONO
        rp.font.size = Pt(9)
        rp.font.color.rgb = TEXTO_CHROME
    tb = slide.shapes.add_textbox(pulg(x + 0.12), pulg(y + chrome + 0.13),
                                  pulg(w - 0.24), pulg(alto - chrome - 0.26))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linea in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(tamano * 1.35)
        r = p.add_run()
        r.text = linea
        r.font.name = MONO
        r.font.size = Pt(tamano)
        if resaltar and any(k in linea for k in resaltar):
            r.font.color.rgb = ROJO_PANEL
            r.font.bold = True
        elif linea.strip().startswith("$"):
            r.font.color.rgb = VERDE_PANEL
        elif "OK" in linea or "passed" in linea or "BUILD SUCCESS" in linea:
            r.font.color.rgb = VERDE_PANEL
        else:
            r.font.color.rgb = TEXTO_PANEL
    return panel


def log_alto(lineas, tamano=9, titulo=None):
    """Alto que ocupara log() con esos parametros, para posicionar lo que sigue."""
    return len(lineas) * (tamano * 1.35 / 72) + 0.26 + (0.24 if titulo else 0.0)


def titulo(s, texto, kicker=None):
    """Titulo + kicker estandar del lienzo claro (diagramas y slides de log)."""
    caja(s, MARGEN, 0.75, 8.9, 0.45, texto, tamano=25, negrita=True)
    if kicker:
        caja(s, MARGEN, 1.22, 8.9, 0.3, kicker, tamano=12, color=TINTA_SUAVE)


def divisor(s, num, lineas, kicker):
    """Divisor de seccion: titulo 34pt, numero de seccion sobre el blob azul
    y kicker bajo el titulo (zona gris, gris permitido)."""
    t = buscar(s, "Texto")
    t.width = pulg(3.9)
    reemplazar(t, lineas, tamano=34)
    caja(s, 0.9, 1.9, 2.2, 1.5, num, tamano=90, negrita=True, color=BLANCO)
    kx = t.left / 914400
    ky = t.top / 914400 + t.height / 914400 + 0.12
    caja(s, kx, ky, 3.9, 0.4, kicker, tamano=12, color=TINTA_SUAVE)


def statement(s, lineas, fuente):
    """Statement de template: cuerpo 27pt, caption de fuente en posicion fija
    y comilla gigante tono-sobre-tono como motivo recurrente (glifo, no franja)."""
    caja(s, 0.7, 0.7, 1.8, 1.7, "“", tamano=130, negrita=True, color=AZUL_OSCURO)
    t = buscar(s, "Simple statement")
    if len(lineas) > 4:
        # Cinco lineas a 27pt: se sube el placeholder para no pisar el caption.
        t.top = pulg(1.7)
        t.height = pulg(2.6)
    reemplazar(t, lineas, tamano=27)
    caja(s, 0.8, 4.4, 4.0, 0.4, fuente, tamano=11, color=BLANCO)


def notas(slide, texto):
    slide.notes_slide.notes_text_frame.text = texto
